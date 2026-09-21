# -*- coding: utf-8 -*-
"""office.py —— 把「教案 / PPT 大纲」落成**真实可打开的** docx / pptx 文件。

设计取舍：
* docx 只用标准库 ``zipfile`` 手写 OOXML 包，**零新增依赖**。
* pptx 有两条路（v7.18 起）：
  - ``_build_pptx_rich()``：**装了 python-pptx 时的默认路径**，按「寻径」视觉规范排版
    （品牌蓝封面页、要点卡片、讲法提示条、页脚页码），这才是给老师看的课件；
  - ``_build_pptx_plain()``：**零依赖兜底**，只用标准库手写 OOXML，一个版式（标题 + 内容）。
    没装 python-pptx（或 rich 版中途出错）时自动回退，保证断网、缺依赖也能出文件。
* 生成的包必须自洽：每个 ``Override PartName`` 都要有实体，每个关系目标都要能解析。
  这一点由 ``validate_package()`` 在写盘前自检，避免"文件生成了但打不开"。
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from math import ceil
from pathlib import Path
from typing import Iterable, Sequence

# 可选依赖：python-pptx 只用于「精美课件」那条路，没装也能跑（回退零依赖版式）。
try:  # pragma: no cover - 取决于运行环境
    from pptx import Presentation as _PptxPresentation
    from pptx.dml.color import RGBColor as _RGBColor
    from pptx.enum.shapes import MSO_SHAPE as _SHAPE
    from pptx.enum.text import MSO_ANCHOR as _ANCHOR, PP_ALIGN as _ALIGN
    from pptx.util import Inches as _Inches, Pt as _Pt

    _PPTX_OK = True
except Exception:  # pragma: no cover - 没装就走兜底
    _PPTX_OK = False

# ================================================================ 公共
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _esc(text: object) -> str:
    """XML 转义 + 剔除 OOXML 不允许的控制字符（粘贴来的文本经常带）。"""
    raw = _ILLEGAL.sub("", str(text if text is not None else ""))
    return (
        raw.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&apos;")
    )


def safe_filename(name: str, ext: str = "") -> str:
    """把标题变成可安全落盘的文件名（去掉 Windows 保留字符）。"""
    stem = re.sub(r'[\\/:*?"<>|\r\n\t]+', "", str(name or "")).strip(" .")
    stem = stem[:60] or "未命名"
    ext = ext if ext.startswith(".") or not ext else f".{ext}"
    return f"{stem}{ext}"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")


# ================================================================ OOXML 骨架
_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_THEME = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="{_A}" name="LearnBuddy">
<a:themeElements>
<a:clrScheme name="LearnBuddy">
<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
<a:dk2><a:srgbClr val="1F2937"/></a:dk2>
<a:lt2><a:srgbClr val="F3F4F6"/></a:lt2>
<a:accent1><a:srgbClr val="2563EB"/></a:accent1>
<a:accent2><a:srgbClr val="0EA5E9"/></a:accent2>
<a:accent3><a:srgbClr val="10B981"/></a:accent3>
<a:accent4><a:srgbClr val="F59E0B"/></a:accent4>
<a:accent5><a:srgbClr val="EF4444"/></a:accent5>
<a:accent6><a:srgbClr val="8B5CF6"/></a:accent6>
<a:hlink><a:srgbClr val="2563EB"/></a:hlink>
<a:folHlink><a:srgbClr val="7C3AED"/></a:folHlink>
</a:clrScheme>
<a:fontScheme name="LearnBuddy">
<a:majorFont><a:latin typeface="Microsoft YaHei"/><a:ea typeface="Microsoft YaHei"/><a:cs typeface=""/></a:majorFont>
<a:minorFont><a:latin typeface="Microsoft YaHei"/><a:ea typeface="Microsoft YaHei"/><a:cs typeface=""/></a:minorFont>
</a:fontScheme>
<a:fmtScheme name="LearnBuddy">
<a:fillStyleLst>
<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
<a:gradFill rotWithShape="1"><a:gsLst>
<a:gs pos="0"><a:schemeClr val="phClr"><a:lumMod val="110000"/><a:satMod val="105000"/></a:schemeClr></a:gs>
<a:gs pos="100000"><a:schemeClr val="phClr"><a:lumMod val="105000"/><a:satMod val="105000"/></a:schemeClr></a:gs>
</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
<a:gradFill rotWithShape="1"><a:gsLst>
<a:gs pos="0"><a:schemeClr val="phClr"><a:lumMod val="115000"/><a:satMod val="105000"/></a:schemeClr></a:gs>
<a:gs pos="100000"><a:schemeClr val="phClr"><a:lumMod val="107000"/><a:satMod val="105000"/></a:schemeClr></a:gs>
</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
</a:fillStyleLst>
<a:lnStyleLst>
<a:ln w="6350" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
<a:ln w="12700" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
<a:ln w="19050" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
</a:lnStyleLst>
<a:effectStyleLst>
<a:effectStyle><a:effectLst/></a:effectStyle>
<a:effectStyle><a:effectLst/></a:effectStyle>
<a:effectStyle><a:effectLst><a:outerShdw blurRad="57150" dist="19050" dir="5400000" algn="ctr" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="63000"/></a:srgbClr></a:outerShdw></a:effectLst></a:effectStyle>
</a:effectStyleLst>
<a:bgFillStyleLst>
<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
<a:solidFill><a:schemeClr val="phClr"><a:tint val="95000"/><a:satMod val="170000"/></a:schemeClr></a:solidFill>
<a:gradFill rotWithShape="1"><a:gsLst>
<a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="93000"/><a:satMod val="150000"/><a:shade val="98000"/><a:lumMod val="102000"/></a:schemeClr></a:gs>
<a:gs pos="100000"><a:schemeClr val="phClr"><a:tint val="98000"/><a:satMod val="130000"/><a:shade val="90000"/><a:lumMod val="103000"/></a:schemeClr></a:gs>
</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
</a:bgFillStyleLst>
</a:fmtScheme>
</a:themeElements>
<a:objectDefaults/><a:extraClrSchemeLst/>
</a:theme>"""

_SPTREE_EMPTY = (
    '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
    '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
)

_SLIDE_MASTER = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldMaster xmlns:a="{_A}" xmlns:r="{_R}" xmlns:p="{_P}">
<p:cSld><p:spTree>{_SPTREE_EMPTY}</p:spTree></p:cSld>
<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>
<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst>
<p:txStyles>
<p:titleStyle><a:lvl1pPr algn="l"><a:defRPr sz="3200" b="1"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill><a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/></a:defRPr></a:lvl1pPr></p:titleStyle>
<p:bodyStyle><a:lvl1pPr marL="342900" indent="-342900"><a:defRPr sz="1800"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/></a:defRPr><a:buChar char="•"/></a:lvl1pPr></p:bodyStyle>
<p:otherStyle><a:lvl1pPr><a:defRPr sz="1800"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/></a:defRPr></a:lvl1pPr></p:otherStyle>
</p:txStyles>
</p:sldMaster>"""

_SLIDE_LAYOUT = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldLayout xmlns:a="{_A}" xmlns:r="{_R}" xmlns:p="{_P}" type="obj" preserve="1">
<p:cSld name="标题和内容"><p:spTree>{_SPTREE_EMPTY}
<p:sp><p:nvSpPr><p:cNvPr id="2" name="Title Placeholder 1"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="838200" y="365125"/><a:ext cx="10515600" cy="1325563"/></a:xfrm></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>
<p:sp><p:nvSpPr><p:cNvPr id="3" name="Content Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph idx="1"/></p:nvPr></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="838200" y="1825625"/><a:ext cx="10515600" cy="4351338"/></a:xfrm></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>
</p:spTree></p:cSld>
<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sldLayout>"""


def _paragraph(text: str, level: int = 0) -> str:
    """内容占位符里的一段（带项目符号层级）。"""
    lvl = max(0, min(4, int(level)))
    return (
        f'<a:p><a:pPr lvl="{lvl}"/>'
        f'<a:r><a:rPr lang="zh-CN" dirty="0"/><a:t>{_esc(text)}</a:t></a:r></a:p>'
    )


def _slide_xml(title: str, bullets: Sequence[str], note: str = "") -> str:
    lines = [_paragraph(b, 0) for b in bullets if str(b).strip()]
    if note:
        lines.append(_paragraph(f"〔讲解提示〕{note}", 0))
    if not lines:
        lines.append(_paragraph("（本页无要点）", 0))
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="{_A}" xmlns:r="{_R}" xmlns:p="{_P}">
<p:cSld><p:spTree>{_SPTREE_EMPTY}
<p:sp><p:nvSpPr><p:cNvPr id="2" name="Title 1"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="838200" y="365125"/><a:ext cx="10515600" cy="1325563"/></a:xfrm></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/>{_paragraph(title)}</p:txBody></p:sp>
<p:sp><p:nvSpPr><p:cNvPr id="3" name="Content 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph idx="1"/></p:nvPr></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="838200" y="1825625"/><a:ext cx="10515600" cy="4351338"/></a:xfrm></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/>{''.join(lines)}</p:txBody></p:sp>
</p:spTree></p:cSld>
<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>"""


def _rel_xml(pairs: Sequence[tuple[str, str, str]]) -> str:
    body = "".join(
        f'<Relationship Id="{rid}" Type="{rtype}" Target="{target}"/>'
        for rid, rtype, target in pairs
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{_RELS}">{body}</Relationships>'
    )


def _core_xml(title: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
        ' xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"'
        ' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"<dc:title>{_esc(title)}</dc:title>"
        "<dc:creator>寻径教育 PathFinder · LearnBuddy</dc:creator>"
        "<cp:lastModifiedBy>寻径教育 PathFinder · LearnBuddy</cp:lastModifiedBy>"
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{_now()}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{_now()}</dcterms:modified>'
        "</cp:coreProperties>"
    )


# ================================================================ 自检
def validate_package(parts: dict[str, bytes]) -> list[str]:
    """检查包自洽性：Override 有实体、关系目标可解析、XML 可解析。返回问题列表。"""
    import xml.dom.minidom as minidom

    problems: list[str] = []
    names = set(parts)

    for name, blob in parts.items():
        if not name.endswith((".xml", ".rels")):
            continue
        try:
            minidom.parseString(blob)
        except Exception as exc:  # pragma: no cover - 自检路径
            problems.append(f"{name} XML 无法解析：{exc}")

    root = parts.get("[Content_Types].xml")
    if root is None:
        problems.append("缺少 [Content_Types].xml")
    else:
        text = root.decode("utf-8")
        for part in re.findall(r'PartName="([^"]+)"', text):
            if part.lstrip("/") not in names:
                problems.append(f"Content_Types 引用了不存在的部件：{part}")

    for name in names:
        if not name.endswith(".rels"):
            continue
        base = name.rsplit("/", 1)[0]
        base = base[:-5] if base.endswith("_rels") else ""
        text = parts[name].decode("utf-8")
        for target in re.findall(r'Target="([^"]+)"', text):
            if target.startswith(("http://", "https://")):
                continue
            if target.startswith("/"):
                resolved = target.lstrip("/")
            else:
                resolved = f"{base}/{target}" if base else target
            # 归一化 ../
            segs: list[str] = []
            for seg in resolved.split("/"):
                if seg == "..":
                    if segs:
                        segs.pop()
                elif seg not in ("", "."):
                    segs.append(seg)
            if "/".join(segs) not in names:
                problems.append(f"{name} 的关系目标缺失：{target}")
    return problems


# ================================================================ PPTX（精美版 · python-pptx）
# 视觉规范照抄项目根 ppt制作提示词.txt：白底大量留白，主色宝蓝、辅色青绿、点缀黄绿，
# 大圆角实色卡片承载要点，统一页脚「左下品牌 + 右下页码」。
_BRAND = "4353FF"      # 宝蓝：标题强调、卡片竖条
_TEAL = "17C9A8"       # 青绿：交替卡片的次级色
_LIME = "D6F94B"       # 荧光黄绿：封面小圆点、讲法提示条
_INK = "111111"
_GRAY = "555555"
_MUTE = "9AA0AE"
_CARD_A = "F4F6FF"     # 极淡蓝卡
_CARD_B = "EEFBF7"     # 极淡青卡
_NOTE_BG = "F8FEE3"    # 极淡黄绿
_FONT = "Microsoft YaHei"

# 16:9 画布（英寸）：左/右边距 0.75，内容宽 11.833
_PW, _PH = 13.333, 7.5
_MX = 0.75
_MW = _PW - _MX * 2


def _style(font, size: float, color: str, bold: bool = False) -> None:
    """统一字形：字号 / 颜色 / 粗体 / 字体名（中西文都要指定，否则汉字会掉回宋体）。"""
    font.size = _Pt(size)
    font.bold = bold
    font.name = _FONT
    font.color.rgb = _RGBColor.from_string(color)
    try:  # 补齐 a:ea，PowerPoint 才知道汉字用哪个字体
        from pptx.oxml.ns import qn

        rPr = getattr(font, "_rPr", None)
        if rPr is None:
            return
        for tag in ("a:latin", "a:ea", "a:cs"):
            found = rPr.find(qn(tag))
            if found is not None:
                found.set("typeface", _FONT)
        if rPr.find(qn("a:ea")) is None:
            latin = rPr.find(qn("a:latin"))
            ea = rPr.makeelement(qn("a:ea"), {"typeface": _FONT})
            latin.addnext(ea) if latin is not None else rPr.append(ea)
    except Exception:  # pragma: no cover - 字体细节不影响出文件
        pass


def _textbox(slide, x: float, y: float, w: float, h: float, text: str,
             size: float = 16, color: str = _INK, bold: bool = False,
             align=None, middle: bool = False, line: float = 1.0):
    """放一段文字。返回文本框，方便调用方继续微调。"""
    box = slide.shapes.add_textbox(_Inches(x), _Inches(y), _Inches(w), _Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR.MIDDLE if middle else _ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = align or _ALIGN.LEFT
    para.line_spacing = line
    run = para.add_run()
    run.text = text
    _style(run.font, size, color, bold)
    return box


def _lines(slide, x: float, y: float, w: float, h: float, items: Sequence[str],
           size: float = 14, color: str = _GRAY, bullet: str = "· ", line: float = 1.45):
    """多行文字（比 add_textbox 多一个「按行分段」的能力，用于封面要点）。"""
    box = slide.shapes.add_textbox(_Inches(x), _Inches(y), _Inches(w), _Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for index, item in enumerate(items):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.line_spacing = line
        run = para.add_run()
        run.text = f"{bullet}{item}"
        _style(run.font, size, color)
    return box


def _rect(slide, shape, x: float, y: float, w: float, h: float, color: str,
          round_rate: float = 0.0):
    """一个纯色块（矩形 / 圆角矩形），无描边无阴影。"""
    sp = slide.shapes.add_shape(shape, _Inches(x), _Inches(y), _Inches(w), _Inches(h))
    sp.shadow.inherit = False
    sp.line.fill.background()
    sp.fill.solid()
    sp.fill.fore_color.rgb = _RGBColor.from_string(color)
    if round_rate and len(sp.adjustments):
        sp.adjustments[0] = round_rate
    return sp


def _cover_slide(slide, title: str, subtitle: str, bullets: Sequence[str], total: int) -> None:
    """封面页：品牌细条 + 标签 + 超大主标题 + 副标题 + 要点 + 底部信息。"""
    _rect(slide, _SHAPE.RECTANGLE, 0, 0, _PW, 0.1, _BRAND)
    _textbox(slide, _MX, 0.34, 6, 0.3, "PathFinder · 寻径教育", 9.5, _MUTE)

    # 黄绿小圆点 + 全大写引导词（照规范里的「小节标签」样式）
    _rect(slide, _SHAPE.OVAL, _MX, 2.12, 0.14, 0.14, _LIME)
    _textbox(slide, _MX + 0.3, 2.02, 8, 0.34, "ONE-CLICK LESSON PREP", 11, _MUTE)

    size = 40 if len(title) <= 12 else 34 if len(title) <= 20 else 28 if len(title) <= 30 else 24
    _textbox(slide, _MX, 2.5, _MW - 1.6, 1.85, title, size, _BRAND, bold=True, line=1.15)
    if subtitle:
        _textbox(slide, _MX, 4.5, _MW - 1.6, 0.5, subtitle, 17, _GRAY)
    if bullets:
        _lines(slide, _MX, 5.25, _MW - 2.4, 1.1, list(bullets)[:3], 14, _GRAY)

    _textbox(slide, _MX, 6.6, _MW, 0.3, datetime.now().strftime("%Y-%m-%d") + " · 一键备课生成", 11, _MUTE)
    _rect(slide, _SHAPE.OVAL, _PW - 1.55, 6.05, 0.52, 0.52, _LIME)
    _footer(slide, 1, total)


def _content_slide(slide, index: int, total: int, title: str,
                   bullets: Sequence[str], note: str) -> None:
    """内容页：左侧竖条大标题 + 要点卡片 + 讲法提示条 + 页脚页码。"""
    _rect(slide, _SHAPE.RECTANGLE, 0, 0, _PW, 0.1, _BRAND)
    _textbox(slide, _MX, 0.34, 6, 0.3, "PathFinder · 寻径教育", 9.5, _MUTE)

    _rect(slide, _SHAPE.RECTANGLE, _MX, 1.05, 0.075, 0.92, _BRAND)
    size = 28 if len(title) <= 14 else 24 if len(title) <= 22 else 20 if len(title) <= 32 else 18
    _textbox(slide, _MX + 0.28, 0.92, _MW - 0.28, 1.2, title, size, _INK, bold=True,
             line=1.15, middle=True)

    items = [str(b).strip() for b in bullets if str(b).strip()]
    note = str(note or "").strip()
    top, bottom = 2.35, 6.05 if note else 6.6
    avail = bottom - top
    gap = 0.2
    cols = 1 if len(items) <= 1 else 2
    rows = max(1, ceil(len(items) / cols))
    cap = 1.5 if cols == 2 else 2.6
    card_h = min((avail - gap * (rows - 1)) / rows, cap)
    start_y = top + (avail - (rows * card_h + gap * (rows - 1))) / 2
    card_w = (_MW - gap * (cols - 1)) / cols

    if not items:
        _textbox(slide, _MX, start_y + 0.4, _MW, 0.5, "（本页无要点）", 14, _MUTE)
    for i, item in enumerate(items):
        row, col = divmod(i, cols)
        cx = _MX + col * (card_w + gap)
        cy = start_y + row * (card_h + gap)
        # 卡片底色与左侧竖条按列交替（蓝 / 青），有节奏但不花
        tint, bar = (_CARD_A, _BRAND) if i % 2 == 0 else (_CARD_B, _TEAL)
        _rect(slide, _SHAPE.ROUNDED_RECTANGLE, cx, cy, card_w, card_h, tint, round_rate=0.05)
        _rect(slide, _SHAPE.RECTANGLE, cx, cy + 0.18, 0.055, card_h - 0.36, bar)
        fsize = 16 if len(item) <= 26 else 14 if len(item) <= 50 else 12
        _textbox(slide, cx + 0.3, cy + 0.14, card_w - 0.55, card_h - 0.28, item,
                 fsize, _INK, middle=True, line=1.3)

    if note:
        _rect(slide, _SHAPE.ROUNDED_RECTANGLE, _MX, 6.2, _MW, 0.62, _NOTE_BG, round_rate=0.12)
        _rect(slide, _SHAPE.RECTANGLE, _MX, 6.2, 0.055, 0.62, _LIME)
        shown = note if len(note) <= 96 else note[:95] + "…"
        _textbox(slide, _MX + 0.28, 6.2, _MW - 0.5, 0.62, "讲法提示：" + shown,
                 12, _GRAY, middle=True)

    _footer(slide, index, total)


def _footer(slide, index: int, total: int) -> None:
    """统一页脚：左下品牌，右下页码。"""
    _textbox(slide, _MX, 6.98, 6, 0.28, "PathFinder · 寻径教育", 9.5, _MUTE)
    _textbox(slide, _PW - _MX - 3, 6.98, 3, 0.28, f"{index:02d} / {total:02d}",
             9.5, _MUTE, align=_ALIGN.RIGHT)


def _build_pptx_rich(title: str, slides: list[dict], subtitle: str = "") -> bytes:
    """用 python-pptx 排出「能直接拿去上课」的课件。"""
    prs = _PptxPresentation()
    prs.slide_width = _Inches(_PW)
    prs.slide_height = _Inches(_PH)
    prs.core_properties.title = title
    prs.core_properties.author = "寻径教育 PathFinder · LearnBuddy"
    blank = prs.slide_layouts[6]  # 空白版式：所有元素自己画，不继承模板占位符

    total = len(slides)
    for i, slide in enumerate(slides, start=1):
        sp = prs.slides.add_slide(blank)
        if i == 1:
            _cover_slide(sp, slide["title"], subtitle, slide["bullets"], total)
        else:
            _content_slide(sp, i, total, slide["title"], slide["bullets"], slide["note"])

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ================================================================ PPTX（零依赖兜底）
def _build_pptx_plain(title: str, slides: Iterable[dict], subtitle: str = "") -> bytes:
    """只用标准库手写 OOXML：一个版式（标题 + 内容），够用且稳。"""
    slides = [
        {
            "title": str(s.get("title") or f"第 {i + 1} 页"),
            "bullets": [str(b) for b in (s.get("bullets") or []) if str(b).strip()],
            "note": str(s.get("note") or ""),
        }
        for i, s in enumerate(slides)
    ] or [{"title": title, "bullets": [], "note": ""}]

    parts: dict[str, bytes] = {}
    put = lambda n, t: parts.__setitem__(n, t.encode("utf-8"))  # noqa: E731

    n = len(slides)
    # --- 内容类型
    overrides = [
        '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>',
        '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>',
        '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>',
        '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>',
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>',
    ]
    for i in range(1, n + 1):
        overrides.append(
            f'<Override PartName="/ppt/slides/slide{i}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        )
    put(
        "[Content_Types].xml",
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        + "".join(overrides)
        + "</Types>",
    )

    put(
        "_rels/.rels",
        _rel_xml(
            [
                ("rId1", f"{_R}/officeDocument", "ppt/presentation.xml"),
                ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
            ]
        ),
    )
    put("docProps/core.xml", _core_xml(title))

    # --- 演示文稿主件
    sld_ids = "".join(f'<p:sldId id="{255 + i}" r:id="rId{i + 1}"/>' for i in range(1, n + 1))
    put(
        "ppt/presentation.xml",
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:presentation xmlns:a="{_A}" xmlns:r="{_R}" xmlns:p="{_P}" saveSubsetFonts="1">'
        '<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
        f"<p:sldIdLst>{sld_ids}</p:sldIdLst>"
        '<p:sldSz cx="12192000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/>'
        "</p:presentation>",
    )

    pres_rels = [("rId1", f"{_R}/slideMaster", "slideMasters/slideMaster1.xml")]
    for i in range(1, n + 1):
        pres_rels.append((f"rId{i + 1}", f"{_R}/slide", f"slides/slide{i}.xml"))
    pres_rels.append((f"rId{n + 2}", f"{_R}/theme", "theme/theme1.xml"))
    put("ppt/_rels/presentation.xml.rels", _rel_xml(pres_rels))

    # --- 母版 / 版式 / 主题
    put("ppt/slideMasters/slideMaster1.xml", _SLIDE_MASTER)
    put(
        "ppt/slideMasters/_rels/slideMaster1.xml.rels",
        _rel_xml(
            [
                ("rId1", f"{_R}/slideLayout", "../slideLayouts/slideLayout1.xml"),
                ("rId2", f"{_R}/theme", "../theme/theme1.xml"),
            ]
        ),
    )
    put("ppt/slideLayouts/slideLayout1.xml", _SLIDE_LAYOUT)
    put(
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels",
        _rel_xml([("rId1", f"{_R}/slideMaster", "../slideMasters/slideMaster1.xml")]),
    )
    put("ppt/theme/theme1.xml", _THEME)

    # --- 幻灯片
    subtitle = str(subtitle or "").strip()
    for i, slide in enumerate(slides, start=1):
        bullets = list(slide["bullets"])
        if i == 1 and subtitle:
            bullets = [subtitle, *bullets]
        put(f"ppt/slides/slide{i}.xml", _slide_xml(slide["title"], bullets, slide["note"]))
        put(
            f"ppt/slides/_rels/slide{i}.xml.rels",
            _rel_xml([("rId1", f"{_R}/slideLayout", "../slideLayouts/slideLayout1.xml")]),
        )

    problems = validate_package(parts)
    if problems:  # pragma: no cover - 只在实现被改坏时触发
        raise RuntimeError("PPTX 包自检未通过：" + "；".join(problems[:5]))

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in parts.items():
            zf.writestr(name, blob)
    return buffer.getvalue()


def _plain(text: str) -> str:
    """把一行「可能是 Markdown」的文本洗成幻灯片上的纯文字。

    检索回来的参考资料片段经常带着 ``# 标题``、``- 要点``、``**加粗**`` 这类记号，
    直接画到卡片上会很脏（甚至出现「# 第3讲 梯度下降与优化」这种页面标题）。
    这里只做最保守的清洗：去行首记号、去加粗星号、压空白，不动正文内容。
    """
    s = str(text or "").strip()
    if not s:
        return ""
    s = re.sub(r"^#{1,6}\s*", "", s)          # ## 标题
    s = re.sub(r"^>\s*", "", s)               # > 引用
    s = re.sub(r"^[-*•·]\s+", "", s)          # - 要点 / * 要点
    s = re.sub(r"^\d+[.、)]\s*", "", s)       # 1. 要点（编号由版式统一给）
    s = s.replace("**", "").replace("`", "")  # 加粗 / 行内代码
    s = re.sub(r"\s+", " ", s).strip(" -—–·")
    return s


def build_pptx(title: str, slides: Iterable[dict], subtitle: str = "") -> bytes:
    """由 PPT 大纲生成 .pptx 字节流（优先精美排版，缺依赖时回退零依赖版式）。

    ``slides`` 每项：``{"title": str, "bullets": [str], "note": str}``。
    """
    clean = [
        {
            "title": _plain(str(s.get("title") or f"第 {i + 1} 页")),
            "bullets": [_plain(str(b)) for b in (s.get("bullets") or [])
                        if _plain(str(b))],
            "note": _plain(str(s.get("note") or "")),
        }
        for i, s in enumerate(slides)
    ] or [{"title": _plain(title), "bullets": [], "note": ""}]

    if _PPTX_OK:
        try:
            return _build_pptx_rich(title, clean, subtitle)
        except Exception:  # pragma: no cover - 兜底路径
            pass
    return _build_pptx_plain(title, clean, subtitle)


# 预览时要剔掉的「版面装饰」：页脚品牌、页码、封面日期行、小节标签
_DECOR_TEXT = ("PathFinder · 寻径教育", "ONE-CLICK LESSON PREP")
_PAGE_NUM = re.compile(r"^\d{1,3}\s*/\s*\d{1,3}$")
_DATE_LINE = re.compile(r"^\d{4}-\d{2}-\d{2}.*一键备课")
_NOTE_PREFIX = "讲法提示："


def _is_decor(line: str) -> bool:
    """这行是版面装饰而不是内容吗？预览只还原内容，装饰交给真正的放映。"""
    return (line in _DECOR_TEXT
            or bool(_PAGE_NUM.match(line))
            or bool(_DATE_LINE.match(line)))


def read_pptx(path: str | Path) -> dict:
    """把 .pptx 读回「文字骨架」，供页面直接翻页预览。

    不还原图片、渐变这些视觉元素，只还原每页的**标题 / 要点 / 备注**——
    老师课前扫一眼内容够不够、顺序对不对，比下载再打开快得多。

    返回 ``{"title": str, "pages": [{"title", "bullets", "note"}]}``；
    没装 python-pptx、文件损坏或压根不是课件时 ``pages`` 为空，前端提示改用「打开」。
    """
    out: dict = {"title": "", "pages": []}
    if not _PPTX_OK:
        return out
    try:
        prs = _PptxPresentation(str(path))
    except Exception:  # pragma: no cover - 坏文件 / 非 OOXML
        return out
    try:
        out["title"] = str(prs.core_properties.title or "").strip()
    except Exception:
        out["title"] = ""
    for sl in prs.slides:
        blocks: list[tuple[int, list[str]]] = []
        try:
            shapes = list(sl.shapes)
        except Exception:
            continue
        for sh in shapes:
            if not getattr(sh, "has_text_frame", False):
                continue          # 纯图形（色块 / 装饰条）没有文字，跳过
            lines = [
                re.sub(r"\s+", " ", p.text).strip()
                for p in sh.text_frame.paragraphs
            ]
            lines = [x for x in lines if x and not _is_decor(x)]
            if lines:
                # 按纵向位置排：先出现的在上面，标题才不会被页脚抢走
                blocks.append((int(getattr(sh, "top", 0) or 0), lines))
        if not blocks:            # 整页没有文字（纯图页）也占位，页码才对得上
            out["pages"].append({"title": "", "bullets": [], "note": ""})
            continue
        blocks.sort(key=lambda b: b[0])
        flat: list[str] = []
        for _, lines in blocks:
            flat.extend(lines)
        title = flat[0]
        rest: list[str] = flat[1:]
        note = ""
        try:
            if sl.has_notes_slide:
                note = re.sub(r"\s+", " ", sl.notes_slide.notes_text_frame.text).strip()
        except Exception:
            note = ""
        if not note:              # 自己生成的课件把「讲法提示」画在页底条里，这里认回来
            for line in rest:
                if line.startswith(_NOTE_PREFIX):
                    note = line[len(_NOTE_PREFIX):].strip()
                    rest.remove(line)
                    break
        out["pages"].append({"title": title, "bullets": [b for b in rest if b][:8], "note": note})
    return out


# ================================================================ DOCX
def _w_para(text: str, style: str = "", bold: bool = False, size: int = 0,
            bullet: bool = False) -> str:
    props: list[str] = []
    if style:
        props.append(f'<w:pStyle w:val="{style}"/>')
    if bullet:
        props.append('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
    run_props = ""
    if bold or size:
        bits = "<w:rPr>"
        if bold:
            bits += "<w:b/>"
        if size:
            bits += f'<w:sz w:val="{int(size)}"/><w:szCs w:val="{int(size)}"/>'
        bits += "</w:rPr>"
        run_props = bits
    ppr = f"<w:pPr>{''.join(props)}</w:pPr>" if props else ""
    return f'<w:p>{ppr}<w:r>{run_props}<w:t xml:space="preserve">{_esc(text)}</w:t></w:r></w:p>'


def build_docx(title: str, blocks: Iterable[Sequence[str]] | Iterable[str]) -> bytes:
    """由「段落块」生成 .docx 字节流。

    ``blocks`` 每项可以是：
    * ``str`` —— 普通段落；
    * ``(文本, 角色)`` —— 角色 ∈ {h1,h2,h3,body,bullet,meta}。
    """
    paras: list[str] = [_w_para(title, size=36, bold=True)]

    for block in blocks:
        if isinstance(block, str):
            paras.append(_w_para(block))
            continue
        text = str(block[0]) if block else ""
        role = str(block[1]) if len(block) > 1 else "body"
        if not text.strip() and role != "body":
            continue
        if role == "h1":
            paras.append(_w_para(text, size=30, bold=True))
        elif role == "h2":
            paras.append(_w_para(text, size=26, bold=True))
        elif role == "h3":
            paras.append(_w_para(text, size=22, bold=True))
        elif role == "bullet":
            paras.append(_w_para(text, bullet=True))
        elif role == "meta":
            paras.append(_w_para(text, size=18))
        else:
            paras.append(_w_para(text))

    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}"><w:body>'
        + "".join(paras)
        + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
        '<w:pgMar w:top="1418" w:right="1418" w:bottom="1418" w:left="1418"/></w:sectPr>'
        "</w:body></w:document>"
    )

    numbering = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:numbering xmlns:w="{_W}">'
        '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0">'
        '<w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="420" w:hanging="420"/></w:pPr>'
        "</w:lvl></w:abstractNum>"
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        "</w:numbering>"
    )

    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
            '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": _rel_xml(
            [
                ("rId1", f"{_R}/officeDocument", "word/document.xml"),
                ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
            ]
        ),
        "docProps/core.xml": _core_xml(title),
        "word/document.xml": document,
        "word/numbering.xml": numbering,
        "word/_rels/document.xml.rels": _rel_xml(
            [("rId1", f"{_R}/numbering", "numbering.xml")]
        ),
    }
    encoded = {k: v.encode("utf-8") for k, v in parts.items()}

    problems = validate_package(encoded)
    if problems:  # pragma: no cover
        raise RuntimeError("DOCX 包自检未通过：" + "；".join(problems[:5]))

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in encoded.items():
            zf.writestr(name, blob)
    return buffer.getvalue()


# ================================================================ 组合导出
def slides_to_text(outline: dict) -> str:
    """PPT 大纲 → 纯文本（与前端 ``PF.slidesToText`` 同形，存资料库 / 建索引用）。

    pptx 本体是二进制，检索不到里面写了什么，所以入库时同时存这份大纲文本：
    老师按关键词能搜到课件，知识点也按「第 N 页 · 标题」逐条切出来。
    """
    outline = outline or {}
    slides = outline.get("slides") or []
    lines = [f"# {outline.get('title') or 'PPT 大纲'}", "", f"共 {len(slides)} 页", ""]
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            continue
        lines.append(f"## 第 {index} 页 · {str(slide.get('title') or '')}")
        for bullet in slide.get("bullets") or []:
            lines.append(f"- {bullet}")
        if slide.get("note"):
            lines.append(f"> 讲法提示：{slide.get('note')}")
        lines.append("")
    return "\n".join(lines)


def lesson_to_blocks(plan: dict) -> list[tuple[str, str]]:
    """把 ``teaching.lesson_plan()['plan']`` 转成 docx 段落块序列。

    只读教案里**确实存在**的字段：
    ``objectives / key_points / difficulties / outline(step,content,minutes) /
    homework / advice / refs / note``。
    """
    plan = plan or {}
    total = plan.get("total_minutes") or (
        int(plan.get("periods") or 1) * 45
    )
    blocks: list[tuple[str, str]] = [
        (
            f"课程：{plan.get('course') or '未指定'}　课时：{plan.get('periods')} 课时"
            f"（共 {total} 分钟）　内容取向：{plan.get('orientation') or '标准型'}",
            "meta",
        ),
        ("", "body"),
    ]

    def section(index: int, heading: str, items, numbered: bool = False) -> None:
        items = [str(i).strip() for i in (items or []) if str(i).strip()]
        if not items:
            return
        blocks.append((f"{index}、{heading}", "h2"))
        for i, item in enumerate(items, start=1):
            blocks.append((f"{i}. {item}" if numbered else item, "bullet"))

    section("一", "教学目标", plan.get("objectives"))
    section("二", "教学重点", plan.get("key_points"))
    section("三", "教学难点", plan.get("difficulties"))

    outline = plan.get("outline") or []
    if outline:
        blocks.append(("四、课堂流程", "h2"))
        for i, seg in enumerate(outline, start=1):
            if not isinstance(seg, dict):
                continue
            blocks.append((
                f"{i}. {seg.get('step') or ''}（{seg.get('minutes') or 0} 分钟）",
                "h3",
            ))
            if seg.get("content"):
                blocks.append((str(seg["content"]), "body"))

    if plan.get("homework"):
        blocks.append(("五、课后作业", "h2"))
        blocks.append((str(plan["homework"]), "body"))
    if plan.get("advice"):
        blocks.append(("六、教师提示", "h2"))
        blocks.append((str(plan["advice"]), "body"))

    refs = [r for r in (plan.get("refs") or []) if str(r).strip()]
    if refs:
        blocks.append(("参考资料", "h2"))
        blocks.append(("本教案依据知识库中的以下切片生成：" + "、".join(f"[{r}]" for r in refs), "meta"))
    if plan.get("note"):
        blocks.append((f"备注：{plan['note']}", "meta"))
    return blocks


if __name__ == "__main__":  # pragma: no cover - 手动验证入口
    demo = build_pptx("测试", [{"title": "首页", "bullets": ["要点一", "要点二"], "note": "提示"}])
    print("pptx bytes:", len(demo))
    print("docx bytes:", len(build_docx("测试教案", [("一、目标", "h2"), ("理解注意力", "bullet")])))
